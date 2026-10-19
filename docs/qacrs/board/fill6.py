import copy, sys
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

INK = RGBColor(0x1C, 0x23, 0x28); GREEN = RGBColor(0x00, 0x8A, 0x58)
p = Presentation('orig.pptx'); s = p.slides[0]
by = {}
for sh in s.shapes: by.setdefault(sh.name, []).append(sh)

def fill(name, lines, size=15, anchor=MSO_ANCHOR.TOP, ltr=False):
    sh = by[name][0]; tf = sh.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for k in ('lIns','rIns','tIns','bIns'): pass
    # clear
    for para in list(tf.paragraphs)[1:]: para._p.getparent().remove(para._p)
    first = True
    for line in lines:
        para = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        for r in list(para.runs): r._r.getparent().remove(r._r)
        para.alignment = PP_ALIGN.LEFT if ltr else PP_ALIGN.RIGHT
        pPr = para._p.get_or_add_pPr(); pPr.set('rtl', '0' if ltr else '1')
        para.space_after = Pt(3)
        parts = line if isinstance(line, list) else [(line, False)]
        for txt, bold in parts:
            r = para.add_run(); r.text = txt
            f = r.font; f.size = Pt(size); f.bold = bool(bold); f.name = 'Arial'
            f.color.rgb = GREEN if bold == 'g' else INK
            if bold == 'g': f.bold = True
            rPr = r._r.get_or_add_rPr(); rPr.set('lang', 'ar-SA')
            for tag in ('a:cs',):
                el = rPr.find(qn(tag))
                if el is None:
                    el = rPr.makeelement(qn(tag), {}); rPr.append(el)
                el.set('typeface', 'Arial')

def set_plain(name, text, size=None):
    sh = by[name][0]; tf = sh.text_frame
    r = tf.paragraphs[0].runs
    r[0].text = text
    for extra in r[1:]: extra.text = ''
    if size: r[0].font.size = Pt(size)

# ---------- header: English acronym title over Arabic
from pptx.enum.shapes import MSO_SHAPE
WHITE = RGBColor(0xFF,0xFF,0xFF)
t = by['عنوان المشروع'][0]; t.top = Emu(2930000); t.height = Emu(1380000)
tf = t.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
for para in list(tf.paragraphs)[1:]: para._p.getparent().remove(para._p)
p0 = tf.paragraphs[0]
for r in list(p0.runs): r._r.getparent().remove(r._r)
p0.alignment = PP_ALIGN.CENTER; p0._p.get_or_add_pPr().set('rtl','0')
r = p0.add_run(); r.text = 'ACRS'; r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name='Arial'
r = p0.add_run(); r.text = '  Automated Compromise Response System'; r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name='Arial'
p1 = tf.add_paragraph(); p1.alignment = PP_ALIGN.CENTER; p1._p.get_or_add_pPr().set('rtl','1'); p1.space_before = Pt(4)
r = p1.add_run(); r.text = 'نظام الاستجابة الآلية للاختراق'; r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name='Arial'
for sh in by['اسم الطالب']:
    if 'اسم الطالب' in sh.text_frame.text and len(sh.text_frame.text) < 20:
        sh.text_frame.paragraphs[0].runs[0].text = 'حسان عبدالله الينبعاوي'; sh.top = Emu(4420000)
sc = by['المدرسة'][0]; sc.top = Emu(4420000)
for r in sc.text_frame.paragraphs[0].runs[1:]: r.text = ''
sc.text_frame.paragraphs[0].runs[0].text = 'ثانوية الموهوبين التقنية'
tg = s.shapes.add_textbox(Emu(2300000), Emu(4960000), Emu(7900000), Emu(520000))
tf = tg.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
pp = tf.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER; pp._p.get_or_add_pPr().set('rtl','1')
r = pp.add_run(); r.text = 'يوقف المخترق بأصغر خطوة آمنة… والخدمات تبقى شغّالة'; r.font.size = Pt(19); r.font.color.rgb = WHITE; r.font.name='Arial'

hdr = by['ترويسة عنوان المشروع'][0]
s.shapes.add_picture('logo.png', hdr.left + Emu(260000), hdr.top + Emu(520000), Emu(1100000), Emu(1000000))
t.left = hdr.left + Emu(1250000); t.width = hdr.width - Emu(1350000)
# ---------- stretch body to fill the space above the footer
Y0, Y1, NEW1 = 5629275, 14782800, 16180000
K = (NEW1 - Y0) / (Y1 - Y0)
for sh in list(s.shapes):
    if Y0 <= sh.top < 16383000 and sh.name != 'التذييل':
        sh.top = Emu(int(Y0 + (sh.top - Y0) * K)); sh.height = Emu(int(sh.height * K))
# ---------- right wing (plain language)
fill('نص المشكلة', [
    'لما يُخترق جهاز، الحل المعتاد قطع الشبكة كلها… فتتعطّل خدمات سليمة.',
    [('في هجوم WannaCry تعطّلت ', False), ('46 جهة صحية', True), (' ما وصلها الهجوم أصلًا.', False)],
    [('والمخترق ينتقل بين الأجهزة في ', False), ('29 دقيقة', True), (' بالمتوسط، وأسرعها 27 ثانية.', False)],
], size=13)
fill('نص الهدف', [
    'نوقف المخترق بأصغر خطوة ممكنة، ونخلي باقي الخدمات شغّالة،',
    [('بدون كسر أي قاعدة أمان', True), ('، وكل قرار موقّع يقدر أحد يراجعه.', False)],
    [('والقرار أسرع من المخترق: ', False), ('أجزاء من الألف من الثانية', True), ('.', False)],
], size=13)
fill('نص الخلفية العلمية', [
    'النظام يحل القرار مثل لغز: يوازن الخيارات ويختار الأفضل، على كمبيوتر عادي وحاسوب كمّي (IBM)، ويعتمد على قائمة MITRE لطرق المخترقين.',
], size=12)
fill('نص الفرضية أو (التصميم الهندسي)', [
    [('نتوقع أن ACRS يوقف المخترق في ', False), ('80%', True), (' أو أكثر مما يوقفه قطع الشبكة كلها، لكن يعطّل خدمات أقل، و', False), ('0', True), (' قرارات خطيرة.', False)],
    'كتبنا التوقع وحفظناه قبل التجربة.',
], size=13)
fill('نص المتغيرات أو (المعايير الهندسية)', [
    [('اللي نغيّره: ', True), ('طريقة الحماية: بدون حماية، أو قطع كل شي، أو قاعدة بسيطة، أو ACRS', False)],
    [('اللي نقيسه: ', True), ('وقف المخترق؟ كم خدمة تعطّلت؟ القرار آمن؟ الوقت.', False)],
    [('اللي نثبّته: ', True), ('نفس الهجمات ونفس الإعدادات للكل.', False)],
], size=13)

# ---------- centre
by['نص الإجراءات'][0].text_frame.text = ''
bg = by['خلفية الإجراءات'][0]
iw = bg.width - Emu(100000); ih = int(iw * 616 / 1650)
s.shapes.add_picture('pipe.png', bg.left + Emu(50000), bg.top + (bg.height - ih)//2, iw, ih)

fill('نص البيانات', [
    [('جرّبنا كل الطرق على نفس الهجمات:', True)],
    [('الضابطة ', 'g'), ('بدون حماية: المخترق يوصل لأجهزة مهمة من 23 طريق', False)],
    [('أ ', 'g'), ('محاكاة: أكثر من 4,000 هجوم، و894 لحظة تشبه الواقع', False)],
    [('ب ', 'g'), ('كمبيوتر Linux حقيقي: نفّذ القرارات صح 66 من 66', False)],
    [('ج ', 'g'), ('حاسوب كمّي (IBM): لقى أفضل حل 11 من 11', False)],
    [('د ', 'g'), ('83 اختراق حقيقي صار في العالم 2013–2026', False)],
], size=14)
ph = by['نص الرسومات البيانية'][0]; ph.text_frame.text = ''
bg = by['خلفية الرسومات البيانية'][0]
img_w = bg.width - Emu(120000); img_h = int(img_w * 470 / 980)
s.shapes.add_picture('chart.png', bg.left + Emu(60000), bg.top + (bg.height - img_h)//2, img_w, img_h)
fill('نص تحليل البيانات', [
    [('اختبار الضغط: ', True), ('1,000 جهاز ← قرار في أقل من 3 أجزاء من الألف من الثانية · 1,000 تنبيه ← 0.7 ث', False)],
    [('اختبار العطل: ', True), ('354 عطلًا محقونًا ← 0 قرار خطير · 200 إعادة تشغيل ← 0 جهاز رجع', False)],
], size=11, anchor=MSO_ANCHOR.MIDDLE)

# ---------- left wing
def kv(num, label): return [(num + '  ', 'g'), (label, False)]
fill('نص النتائج', [
    kv('91%', 'من الهجمات احتواها تلقائيًا'),
    kv('70', 'خدمة تعطّلت بدل 268'),
    kv('0', 'قرارات خطيرة'),
    kv('11/11', 'أفضل حل كمّي'),
    kv('0', 'جهاز رجع بعد 200 إعادة تشغيل'),
], size=12)
for para in by['نص النتائج'][0].text_frame.paragraphs:
    para.runs[0].font.size = Pt(18); para.space_after = Pt(0)
fill('نص الاستنتاج', [
    'نقدر نوقف المخترق بدون ما نطفي الخدمات، من أول تنبيه.',
    [('مثال HSE (مستشفيات إيرلندا): ', True), ('فصل جهاز واحد كان يكفي، بدل تعطّل 9 خدمات بعد 6 أسابيع.', False)],
], size=12)
fill('نص التطبيقات', [
    'المستشفيات، والبنوك، والجامعات، والمرافق المهمة مثل الكهرباء والماء.',
    'مصمَّم لدعم ضوابط الأمن السيبراني الوطنية ECC-2:2024.',
], size=13)
fill('نص الأعمال المستقبلية', [
    'جهاز ACRS Box صغير يُركّب بالشبكة، ومواجهة المخترق الذكي، وتجربة الحاسوب الكمّي في أرامكو، ونشر بحث علمي.',
], size=13)
fill('نص مصادر البحث', ltr=True, lines=[
    'NAO (2017). WannaCry and the NHS.',
    'PwC (2021). Conti attack on the HSE.',
    'NCA (2024). ECC-2:2024.',
    'MITRE (2024). Attack Flow corpus.',
    'CrowdStrike (2026). Global Threat Report.',
], size=12)

p.save('ACRS_board.pptx'); print('saved')
