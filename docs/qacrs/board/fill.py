import copy, sys
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

INK = RGBColor(0x1C, 0x23, 0x28); GREEN = RGBColor(0x00, 0x8A, 0x58)
p = Presentation('orig.pptx'); s = p.slides[0]
by = {}
for sh in s.shapes: by.setdefault(sh.name, []).append(sh)

def fill(name, lines, size=13, anchor=MSO_ANCHOR.TOP, ltr=False):
    sh = by[name][0]; tf = sh.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for k in ('lIns','rIns','tIns','bIns'): pass
    # clear
    for para in list(tf.paragraphs)[1:]: para._p.getparent().remove(para._p)
    first = True
    for line in lines:
        para = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        for r in list(para.runs): r._r.getparent().remove(r._r)
        para.alignment = PP_ALIGN.LEFT if ltr else PP_ALIGN.RIGHT
        pPr = para._p.get_or_add_pPr(); pPr.set('rtl', '0' if ltr else '1')
        para.space_after = Pt(3)
        parts = line if isinstance(line, list) else [(line, False)]
        for txt, bold in parts:
            r = para.add_run(); r.text = txt
            f = r.font; f.size = Pt(size); f.bold = bool(bold); f.name = 'Arial'
            f.color.rgb = GREEN if bold == 'g' else INK
            if bold == 'g': f.bold = True
            rPr = r._r.get_or_add_rPr(); rPr.set('lang', 'ar-SA')
            for tag in ('a:cs',):
                el = rPr.find(qn(tag))
                if el is None:
                    el = rPr.makeelement(qn(tag), {}); rPr.append(el)
                el.set('typeface', 'Arial')

def set_plain(name, text, size=None):
    sh = by[name][0]; tf = sh.text_frame
    r = tf.paragraphs[0].runs
    r[0].text = text
    for extra in r[1:]: extra.text = ''
    if size: r[0].font.size = Pt(size)

# ---------- header
set_plain('عنوان المشروع', 'ACRS: احتواء الهجوم بأقل إجراء آمن', 29)
for sh in by['اسم الطالب']:
    if 'اسم الطالب' in sh.text_frame.text and len(sh.text_frame.text) < 20:
        sh.text_frame.paragraphs[0].runs[0].text = 'حسان عبدالله الينبعاوي'
for r in by['المدرسة'][0].text_frame.paragraphs[0].runs[1:]: r.text = ''
by['المدرسة'][0].text_frame.paragraphs[0].runs[0].text = 'ثانوية الموهوبين التقنية'

# ---------- right wing (problem side)
fill('نص المشكلة', [
    'العزل الكامل عند الهجوم يوقف خدمات سليمة.',
    [('في WannaCry تعطلت ', False), ('46 جهة صحية', True), (' لم تُصب أصلًا.', False)],
])
fill('نص الهدف', [
    'اختيار أقل إجراء يوقف المهاجم ويُبقي الخدمات تعمل،',
    [('دون كسر أي قاعدة سلامة', True), (' وبقرار قابل للتحقق.', False)],
])
fill('نص الخلفية العلمية', [
    'تحسين ثنائي QUBO يُحل كلاسيكيًا وعلى حاسوب IBM الكمّي (QAOA)، مع MITRE ATT&CK وCVSS.',
], size=12)
fill('نص الفرضية أو (التصميم الهندسي)', [
    [('ACRS يحتوي ', False), ('≥ 80%', True), (' مما يحتويه العزل الكامل، بخدمات متوقفة أقل، و', False), ('0', True), (' قرارات غير آمنة.', False)],
    'كل فرضية مسجّلة في Git قبل التجربة.',
])
fill('نص المتغيرات أو (المعايير الهندسية)', [
    [('المستقل: ', True), ('سياسة القرار (بلا حماية، عزل كامل، قاعدة بسيطة، ACRS).', False)],
    [('التابع: ', True), ('الاحتواء، الخدمات المتوقفة، السلامة، الزمن.', False)],
    [('الضابط: ', True), ('الأوزان، والعتبة 7.', False)],
], size=12)

# ---------- center
fill('نص الإجراءات', [
    [('١ ', 'g'), ('تنبيه بالهجوم (تقنية ATT&CK على جهاز)', False)],
    [('٢ ', 'g'), ('أرضية أمان: مشفَّر ← عزل، خادم حرج ← تجزئة تُبقي خدمته', False)],
    [('٣ ', 'g'), ('مُحسِّن QUBO يختار أقل إجراء لبقية الشبكة', False)],
    [('٤ ', 'g'), ('بوابة سلامة مستقلة: القرار الحرج لموافقة شخصين', False)],
    [('٥ ', 'g'), ('تنفيذ في جدار Linux حقيقي + شهادة موقّعة', False)],
], size=15, anchor=MSO_ANCHOR.MIDDLE)
fill('نص البيانات', [
    [('83', True), (' حادثة حقيقية 2013–2026، منها WannaCry وColonial وHSE', False)],
    [('4,000+', True), (' سيناريو تركيبي و', False), ('894', True), (' واقعي، على شبكات صحية وصناعية ومصرفية وجامعية', False)],
    [('3', True), (' حواسيب IBM كمّية، وجدار Linux حقيقي', False)],
], size=13)
fill('نص تحليل البيانات', [
    'فترات ثقة Wilson، وMcNemar، وWilcoxon (α = 0.05)؛ والخدمات المتوقفة تُحسب بعد موافقة البشر.',
], size=12, anchor=MSO_ANCHOR.MIDDLE)
# chart into the charts frame
ph = by['نص الرسومات البيانية'][0]; ph.text_frame.text = ''
bg = by['خلفية الرسومات البيانية'][0]
img_w = bg.width - Emu(120000); img_h = int(img_w * 400 / 980)
s.shapes.add_picture('chart.png', bg.left + Emu(60000), bg.top + (bg.height - img_h)//2, img_w, img_h)

# ---------- left wing (results side)
fill('نص النتائج', [
    [('91%', 'g'), (' احتواء تلقائي (75 حادثة جديدة)', False)],
    [('70', 'g'), (' خدمة متوقفة مقابل ', False), ('268', True)],
    [('0', 'g'), (' قرارات غير آمنة', False)],
    [('11/11', 'g'), (' حل أمثل على IBM الكمّي', False)],
], size=13)
fill('نص الاستنتاج', [
    'يمكن احتواء الهجوم دون إيقاف الخدمات متى وصل التنبيه.',
    [('في HSE: ', True), ('جهاز واحد عند أول تنبيه بدل 9 خدمات بعد 6 أسابيع.', False)],
], size=12)
fill('نص التطبيقات', [
    'المستشفيات، والبنوك، والجامعات، والبنى الحيوية.',
    'مصمَّم لدعم الضوابط الأساسية للأمن السيبراني ECC-2:2024.',
], size=12)
fill('نص الأعمال المستقبلية', [
    'جهاز ACRS Box (Raspberry Pi 5)، ومقاومة المهاجم الذكي، والتنفيذ على حاسوب أرامكو/Pasqal الكمّي، ونشر ورقة علمية.',
], size=12)
fill('نص مصادر البحث', ltr=True, lines=[
    'NAO (2017). WannaCry and the NHS.',
    'PwC (2021). Conti attack on the HSE.',
    'NCA (2024). ECC-2:2024.',
    'github.com/hasa32054-crypto/ACRS',
], size=11)
p.save('ACRS_board.pptx'); print('saved')
