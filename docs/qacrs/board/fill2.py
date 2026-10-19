import copy, sys
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

INK = RGBColor(0x1C, 0x23, 0x28); GREEN = RGBColor(0x00, 0x8A, 0x58)
p = Presentation('orig.pptx'); s = p.slides[0]
by = {}
for sh in s.shapes: by.setdefault(sh.name, []).append(sh)

def fill(name, lines, size=13, anchor=MSO_ANCHOR.TOP, ltr=False):
    sh = by[name][0]; tf = sh.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for k in ('lIns','rIns','tIns','bIns'): pass
    # clear
    for para in list(tf.paragraphs)[1:]: para._p.getparent().remove(para._p)
    first = True
    for line in lines:
        para = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        for r in list(para.runs): r._r.getparent().remove(r._r)
        para.alignment = PP_ALIGN.LEFT if ltr else PP_ALIGN.RIGHT
        pPr = para._p.get_or_add_pPr(); pPr.set('rtl', '0' if ltr else '1')
        para.space_after = Pt(3)
        parts = line if isinstance(line, list) else [(line, False)]
        for txt, bold in parts:
            r = para.add_run(); r.text = txt
            f = r.font; f.size = Pt(size); f.bold = bool(bold); f.name = 'Arial'
            f.color.rgb = GREEN if bold == 'g' else INK
            if bold == 'g': f.bold = True
            rPr = r._r.get_or_add_rPr(); rPr.set('lang', 'ar-SA')
            for tag in ('a:cs',):
                el = rPr.find(qn(tag))
                if el is None:
                    el = rPr.makeelement(qn(tag), {}); rPr.append(el)
                el.set('typeface', 'Arial')

def set_plain(name, text, size=None):
    sh = by[name][0]; tf = sh.text_frame
    r = tf.paragraphs[0].runs
    r[0].text = text
    for extra in r[1:]: extra.text = ''
    if size: r[0].font.size = Pt(size)

# ---------- header: English acronym title over Arabic
from pptx.enum.shapes import MSO_SHAPE
WHITE = RGBColor(0xFF,0xFF,0xFF)
t = by['عنوان المشروع'][0]; t.top = Emu(2930000); t.height = Emu(1380000)
tf = t.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
for para in list(tf.paragraphs)[1:]: para._p.getparent().remove(para._p)
p0 = tf.paragraphs[0]
for r in list(p0.runs): r._r.getparent().remove(r._r)
p0.alignment = PP_ALIGN.CENTER; p0._p.get_or_add_pPr().set('rtl','0')
r = p0.add_run(); r.text = 'ACRS'; r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name='Arial'
r = p0.add_run(); r.text = '  Automated Compromise Response System'; r.font.size = Pt(22); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name='Arial'
p1 = tf.add_paragraph(); p1.alignment = PP_ALIGN.CENTER; p1._p.get_or_add_pPr().set('rtl','1'); p1.space_before = Pt(4)
r = p1.add_run(); r.text = 'نظام الاستجابة الآلية للاختراق'; r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name='Arial'
for sh in by['اسم الطالب']:
    if 'اسم الطالب' in sh.text_frame.text and len(sh.text_frame.text) < 20:
        sh.text_frame.paragraphs[0].runs[0].text = 'حسان عبدالله الينبعاوي'; sh.top = Emu(4420000)
sc = by['المدرسة'][0]; sc.top = Emu(4420000)
for r in sc.text_frame.paragraphs[0].runs[1:]: r.text = ''
sc.text_frame.paragraphs[0].runs[0].text = 'ثانوية الموهوبين التقنية'
tg = s.shapes.add_textbox(Emu(2300000), Emu(4960000), Emu(7900000), Emu(520000))
tf = tg.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
pp = tf.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER; pp._p.get_or_add_pPr().set('rtl','1')
r = pp.add_run(); r.text = 'يوقف المهاجم بأقل إجراء آمن… ويُبقي الخدمات تعمل'; r.font.size = Pt(19); r.font.color.rgb = WHITE; r.font.name='Arial'

# ---------- right wing
fill('نص المشكلة', [
    'العزل الكامل عند الهجوم يوقف خدمات سليمة.',
    [('في WannaCry تعطلت ', False), ('46 جهة صحية', True), (' لم تُصب أصلًا.', False)],
])
fill('نص الهدف', [
    'اختيار أقل إجراء يوقف المهاجم ويُبقي الخدمات تعمل،',
    [('دون كسر أي قاعدة سلامة', True), (' وبقرار قابل للتحقق.', False)],
])
fill('نص الخلفية العلمية', [
    'تحسين ثنائي QUBO يُحل كلاسيكيًا وعلى حاسوب IBM الكمّي (QAOA)، مع MITRE ATT&CK وCVSS.',
], size=12)
fill('نص الفرضية أو (التصميم الهندسي)', [
    [('ACRS يحتوي ', False), ('≥ 80%', True), (' مما يحتويه العزل الكامل، بخدمات متوقفة أقل، و', False), ('0', True), (' قرارات غير آمنة.', False)],
    'كل فرضية مسجّلة في Git قبل التجربة.',
])
fill('نص المتغيرات أو (المعايير الهندسية)', [
    [('المستقل: ', True), ('سياسة القرار: الضابطة، العزل الكامل، القاعدة البسيطة، ACRS.', False)],
    [('التابع: ', True), ('الاحتواء، الخدمات المتوقفة، السلامة، الزمن.', False)],
    [('الضابط: ', True), ('الأوزان، والعتبة 7، والبذور.', False)],
], size=12)

# ---------- centre: engineering drawing, experimental groups, load notebook
by['نص الإجراءات'][0].text_frame.text = ''
bg = by['خلفية الإجراءات'][0]
iw = bg.width - Emu(100000); ih = int(iw * 500 / 1650)
s.shapes.add_picture('pipe.png', bg.left + Emu(50000), bg.top + (bg.height - ih)//2, iw, ih)

fill('نص البيانات', [
    [('المجموعات التجريبية على نفس الهجمات:', True)],
    [('الضابطة ', 'g'), ('بلا حماية: 23 مسارًا مفتوحًا إلى خوادم حرجة', False)],
    [('أ ', 'g'), ('محاكاة: 4,000+ سيناريو تركيبي و894 واقعي', False)],
    [('ب ', 'g'), ('نواة Linux: نفس القواعد تُختبر حاسوبيًا ثم فعليًا (66/66)', False)],
    [('ج ', 'g'), ('كمّي: 3 حواسيب IBM، الأمثل 11/11', False)],
    [('د ', 'g'), ('الابتكار: 83 حادثة حقيقية 2013–2026', False)],
], size=12)
ph = by['نص الرسومات البيانية'][0]; ph.text_frame.text = ''
bg = by['خلفية الرسومات البيانية'][0]
img_w = bg.width - Emu(120000); img_h = int(img_w * 400 / 980)
s.shapes.add_picture('chart.png', bg.left + Emu(60000), bg.top + (bg.height - img_h)//2, img_w, img_h)
fill('نص تحليل البيانات', [
    [('نوتة الحمل: ', True), ('1,000 جهاز ← 2.3 مللي ثانية · 1,000 تنبيه دفعة ← 0.7 ثانية · Wilson وMcNemar وWilcoxon', False)],
], size=11, anchor=MSO_ANCHOR.MIDDLE)

# ---------- left wing: big numbers
def kv(num, label): return [(num + '  ', 'g'), (label, False)]
fill('نص النتائج', [
    kv('91%', 'احتواء تلقائي'),
    kv('70', 'خدمة متوقفة مقابل 268'),
    kv('0', 'قرارات غير آمنة'),
    kv('11/11', 'حل أمثل على IBM الكمّي'),
], size=12)
for para in by['نص النتائج'][0].text_frame.paragraphs:
    para.runs[0].font.size = Pt(15); para.space_after = Pt(1)
fill('نص الاستنتاج', [
    'يمكن احتواء الهجوم دون إيقاف الخدمات متى وصل التنبيه.',
    [('في HSE: ', True), ('جهاز واحد عند أول تنبيه بدل 9 خدمات بعد 6 أسابيع.', False)],
], size=12)
fill('نص التطبيقات', [
    'المستشفيات، والبنوك، والجامعات، والبنى الحيوية.',
    'مصمَّم لدعم الضوابط الأساسية للأمن السيبراني ECC-2:2024.',
], size=12)
fill('نص الأعمال المستقبلية', [
    'جهاز ACRS Box (Raspberry Pi 5)، ومقاومة المهاجم الذكي، والتنفيذ على حاسوب أرامكو/Pasqal الكمّي، ونشر ورقة علمية.',
], size=12)
fill('نص مصادر البحث', ltr=True, lines=[
    'NAO (2017). WannaCry and the NHS.',
    'PwC (2021). Conti attack on the HSE.',
    'NCA (2024). ECC-2:2024.',
    'MITRE (2024). Attack Flow corpus.',
], size=11)

# ---------- KPI band + QR between the wings and the footer
Y = 14980000; H = 1250000
tiles = [('91%', 'احتواء تلقائي على 75 هجمة حقيقية جديدة', '417FB5'),
         ('70 / 268', 'خدمة متوقفة مقابل العزل الكامل', '0DA4A6'),
         ('0', 'قرارات غير آمنة في كل التجارب', '00A866')]
x_right = 12458700 - 300000; tw = 3180000; gap = 200000
for i, (num, lab, col) in enumerate(tiles):
    x = x_right - (i + 1) * tw - i * gap
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(Y), Emu(tw), Emu(H))
    box.adjustments[0] = 0.12; box.fill.solid(); box.fill.fore_color.rgb = RGBColor.from_string(col); box.line.fill.background()
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    a = tf.paragraphs[0]; a.alignment = PP_ALIGN.CENTER; a._p.get_or_add_pPr().set('rtl','0')
    rr = a.add_run(); rr.text = num; rr.font.size = Pt(34); rr.font.bold = True; rr.font.color.rgb = WHITE; rr.font.name='Arial'
    b2 = tf.add_paragraph(); b2.alignment = PP_ALIGN.CENTER; b2._p.get_or_add_pPr().set('rtl','1')
    rr = b2.add_run(); rr.text = lab; rr.font.size = Pt(13); rr.font.color.rgb = WHITE; rr.font.name='Arial'
qx = 300000; qs = 1020000
s.shapes.add_picture('qr.png', Emu(qx + (1300000-qs)//2), Emu(Y), Emu(qs), Emu(qs))
cap = s.shapes.add_textbox(Emu(qx), Emu(Y + qs + 10000), Emu(1300000), Emu(230000))
cp = cap.text_frame.paragraphs[0]; cp.alignment = PP_ALIGN.CENTER; cp._p.get_or_add_pPr().set('rtl','1')
rr = cp.add_run(); rr.text = 'الكود والنتائج'; rr.font.size = Pt(11); rr.font.bold = True; rr.font.color.rgb = INK; rr.font.name='Arial'
p.save('ACRS_board.pptx'); print('saved')
